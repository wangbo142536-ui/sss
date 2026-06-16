import argparse
import json
import subprocess
import tempfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


SOURCE_FILE = "IMPA.xlsx"
ZH_LANGUAGE = "zh-CN"
EN_LANGUAGE = "en-US"
TARGET_TABLES = {
    "impa_item",
    "impa_item_i18n",
    "data_import_batch",
    "data_import_error",
}


@dataclass(frozen=True)
class ImpaItemRow:
    impa_code: str
    category_code: str
    default_unit_cn: str | None
    default_unit_en: str | None
    specification_cn: str | None
    specification_en: str | None
    source_row_no_cn: int
    source_row_no_en: int


@dataclass(frozen=True)
class ImpaI18nRow:
    impa_code: str
    category_code: str
    language: str
    description: str
    specification: str | None
    unit: str | None
    remark: str | None
    source_sheet: str
    source_row_no: int


@dataclass(frozen=True)
class ImportErrorRow:
    source_sheet: str
    source_row_no: int
    business_key: str
    error_code: str
    error_message: str
    raw_payload: dict


@dataclass(frozen=True)
class ImportPlan:
    items: list[ImpaItemRow]
    i18n_rows: list[ImpaI18nRow]
    errors: list[ImportErrorRow]
    target_tables: set[str]


def normalize(value):
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    return text.replace("_x000D_", "\n")


def load_sheet_rows(workbook, sheet_index):
    sheet = workbook.worksheets[sheet_index]
    header = next(sheet.iter_rows(values_only=True))
    indexes = {str(value).strip(): index for index, value in enumerate(header) if value is not None}
    rows = {}
    errors = []

    for row_no, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not any(value is not None for value in row):
            continue

        impa_code = normalize(row[indexes["IMPA"]])
        category_code = normalize(row[indexes["CODE"]])
        description = normalize(row[indexes["DESCRIPTION"]])
        payload = {
            "IMPA": impa_code,
            "CODE": category_code,
            "DESCRIPTION": description,
            "SPECIFICATION": normalize(row[indexes["SPECIFICATION"]]),
            "UNIT": normalize(row[indexes["UNIT"]]),
            "REMARK": normalize(row[indexes["REMARK"]]),
        }

        if not impa_code or not category_code:
            errors.append(
                ImportErrorRow(
                    source_sheet=sheet.title,
                    source_row_no=row_no,
                    business_key=impa_code or "",
                    error_code="MISSING_KEY",
                    error_message="IMPA or CODE is empty.",
                    raw_payload=payload,
                )
            )
            continue

        if impa_code in rows:
            errors.append(
                ImportErrorRow(
                    source_sheet=sheet.title,
                    source_row_no=row_no,
                    business_key=impa_code,
                    error_code="DUPLICATE_IMPA",
                    error_message="Duplicate IMPA code in source sheet.",
                    raw_payload=payload,
                )
            )
            continue

        if not impa_code.startswith(category_code):
            errors.append(
                ImportErrorRow(
                    source_sheet=sheet.title,
                    source_row_no=row_no,
                    business_key=impa_code,
                    error_code="CODE_PREFIX_MISMATCH",
                    error_message="IMPA code does not start with CODE.",
                    raw_payload=payload,
                )
            )

        rows[impa_code] = {
            "row_no": row_no,
            "sheet": sheet.title,
            **payload,
        }

    return rows, errors


def build_import_plan(workbook_path: Path) -> ImportPlan:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    zh_rows, zh_errors = load_sheet_rows(workbook, 1)
    en_rows, en_errors = load_sheet_rows(workbook, 2)

    items = []
    i18n_rows = []
    errors = [*zh_errors, *en_errors]

    for impa_code in sorted(set(zh_rows) | set(en_rows)):
        zh = zh_rows.get(impa_code)
        en = en_rows.get(impa_code)

        if zh is None or en is None:
            present_sheet = zh or en
            errors.append(
                ImportErrorRow(
                    source_sheet=present_sheet["sheet"],
                    source_row_no=present_sheet["row_no"],
                    business_key=impa_code,
                    error_code="LANGUAGE_ROW_MISMATCH",
                    error_message="IMPA code is not present in both Chinese and English sheets.",
                    raw_payload=present_sheet,
                )
            )
            continue

        items.append(
            ImpaItemRow(
                impa_code=impa_code,
                category_code=zh["CODE"],
                default_unit_cn=zh["UNIT"],
                default_unit_en=en["UNIT"],
                specification_cn=zh["SPECIFICATION"],
                specification_en=en["SPECIFICATION"],
                source_row_no_cn=zh["row_no"],
                source_row_no_en=en["row_no"],
            )
        )

        for language, source in ((ZH_LANGUAGE, zh), (EN_LANGUAGE, en)):
            if not source["DESCRIPTION"]:
                errors.append(
                    ImportErrorRow(
                        source_sheet=source["sheet"],
                        source_row_no=source["row_no"],
                        business_key=impa_code,
                        error_code="MISSING_DESCRIPTION",
                        error_message="DESCRIPTION is empty; i18n row was skipped.",
                        raw_payload=source,
                    )
                )
                continue

            i18n_rows.append(
                ImpaI18nRow(
                    impa_code=impa_code,
                    category_code=source["CODE"],
                    language=language,
                    description=source["DESCRIPTION"],
                    specification=source["SPECIFICATION"],
                    unit=source["UNIT"],
                    remark=source["REMARK"],
                    source_sheet=source["sheet"],
                    source_row_no=source["row_no"],
                )
            )

    return ImportPlan(
        items=items,
        i18n_rows=i18n_rows,
        errors=errors,
        target_tables=TARGET_TABLES,
    )


def sql_value(value):
    if value is None:
        return "NULL"
    if isinstance(value, (int, float)):
        return str(value)
    escaped = str(value).replace("\\", "\\\\").replace("'", "''")
    return f"'{escaped}'"


def write_values_statement(handle, table_name, columns, rows, update_columns=None, batch_size=500):
    for start in range(0, len(rows), batch_size):
        batch = rows[start : start + batch_size]
        handle.write(f"INSERT INTO {table_name} ({', '.join(columns)}) VALUES\n")
        values_sql = []
        for row in batch:
            values_sql.append("(" + ", ".join(sql_value(value) for value in row) + ")")
        handle.write(",\n".join(values_sql))
        if update_columns:
            update_parts = []
            for column in update_columns:
                if column == "updated_at":
                    update_parts.append("updated_at = CURRENT_TIMESTAMP")
                else:
                    update_parts.append(f"{column} = VALUES({column})")
            updates = ", ".join(update_parts)
            handle.write(f"\nON DUPLICATE KEY UPDATE {updates};\n")
        else:
            handle.write(";\n")


def write_import_sql(plan: ImportPlan, output_path: Path):
    output_path.parent.mkdir(parents=True, exist_ok=True)
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    with output_path.open("w", encoding="utf-8", newline="\n") as handle:
        handle.write("USE ship_supply_platform;\n")
        handle.write("START TRANSACTION;\n")
        handle.write(
            "INSERT INTO data_import_batch "
            "(import_type, source_file, source_sheet, total_rows, success_rows, failed_rows, status, started_at, finished_at) "
            f"VALUES ('IMPA', '{SOURCE_FILE}', NULL, {len(plan.items)}, {len(plan.items)}, {len(plan.errors)}, "
            f"'PARTIAL_SUCCESS', '{now}', '{now}');\n"
        )
        handle.write("SET @batch_id = LAST_INSERT_ID();\n")

        write_values_statement(
            handle,
            "impa_item",
            [
                "impa_code",
                "category_code",
                "default_unit_cn",
                "default_unit_en",
                "specification_cn",
                "specification_en",
                "enabled",
                "source_file",
                "source_row_no_cn",
                "source_row_no_en",
            ],
            [
                (
                    row.impa_code,
                    row.category_code,
                    row.default_unit_cn,
                    row.default_unit_en,
                    row.specification_cn,
                    row.specification_en,
                    1,
                    SOURCE_FILE,
                    row.source_row_no_cn,
                    row.source_row_no_en,
                )
                for row in plan.items
            ],
            update_columns=[
                "category_code",
                "default_unit_cn",
                "default_unit_en",
                "specification_cn",
                "specification_en",
                "enabled",
                "source_file",
                "source_row_no_cn",
                "source_row_no_en",
                "updated_at",
            ],
        )

        write_values_statement(
            handle,
            "impa_item_i18n",
            [
                "impa_code",
                "category_code",
                "language",
                "description",
                "specification",
                "unit",
                "remark",
                "source_sheet",
                "source_row_no",
            ],
            [
                (
                    row.impa_code,
                    row.category_code,
                    row.language,
                    row.description,
                    row.specification,
                    row.unit,
                    row.remark,
                    row.source_sheet,
                    row.source_row_no,
                )
                for row in plan.i18n_rows
            ],
            update_columns=[
                "category_code",
                "description",
                "specification",
                "unit",
                "remark",
                "source_sheet",
                "source_row_no",
                "updated_at",
            ],
        )

        if plan.errors:
            write_values_statement(
                handle,
                "data_import_error",
                [
                    "batch_id",
                    "source_sheet",
                    "source_row_no",
                    "business_key",
                    "error_code",
                    "error_message",
                    "raw_payload",
                ],
                [
                    (
                        "@batch_id",
                        error.source_sheet,
                        error.source_row_no,
                        error.business_key,
                        error.error_code,
                        error.error_message,
                        json.dumps(error.raw_payload, ensure_ascii=False),
                    )
                    for error in plan.errors
                ],
            )
        handle.write("COMMIT;\n")


def replace_batch_id_variable(sql_path: Path):
    text = sql_path.read_text(encoding="utf-8")
    text = text.replace("'@batch_id'", "@batch_id")
    sql_path.write_text(text, encoding="utf-8", newline="\n")


def execute_sql(mysql_path: Path, sql_path: Path, user: str, password: str, host: str, port: int):
    command = [
        str(mysql_path),
        f"--host={host}",
        f"--port={port}",
        f"--user={user}",
        f"--password={password}",
        "--default-character-set=utf8mb4",
        "ship_supply_platform",
    ]
    with sql_path.open("rb") as sql_file:
        return subprocess.run(command, stdin=sql_file, check=False, capture_output=True)


def main():
    parser = argparse.ArgumentParser(description="Import IMPA standard item rows into MySQL.")
    parser.add_argument("--workbook", default="tmp/provision_analysis/IMPA.xlsx")
    parser.add_argument("--sql-output", default="tmp/provision_analysis/impa_items_import.sql")
    parser.add_argument("--execute", action="store_true")
    parser.add_argument("--mysql", default=r"C:\Program Files\MySQL\MySQL Server 8.0\bin\mysql.exe")
    parser.add_argument("--host", default="127.0.0.1")
    parser.add_argument("--port", default=3306, type=int)
    parser.add_argument("--user", default="root")
    parser.add_argument("--password", default="root")
    args = parser.parse_args()

    plan = build_import_plan(Path(args.workbook))
    sql_path = Path(args.sql_output)
    write_import_sql(plan, sql_path)
    replace_batch_id_variable(sql_path)

    print(
        json.dumps(
            {
                "items": len(plan.items),
                "i18n_rows": len(plan.i18n_rows),
                "errors": len(plan.errors),
                "target_tables": sorted(plan.target_tables),
                "sql_output": str(sql_path),
            },
            ensure_ascii=False,
        )
    )

    if args.execute:
        result = execute_sql(Path(args.mysql), sql_path, args.user, args.password, args.host, args.port)
        if result.stdout:
            print(result.stdout.decode("utf-8", errors="replace"))
        if result.stderr:
            print(result.stderr.decode("utf-8", errors="replace"))
        raise SystemExit(result.returncode)


if __name__ == "__main__":
    main()
